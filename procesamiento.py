"""Lógica de negocio: adaptación a Python de las macros VBA de reportes de horarios.

Flujo principal:
    leer_reporte_inspector()  -> bloques de 5 filas por empleado a tabla plana (incluye TOLERANCIAS)
    construir_datos()         -> tabla de horarios con TURNO (activos + en firma)
    evaluar_turnos()          -> reglas de turno (macros 5 y 6)
    leer_csv_mysql()          -> export de MySQL (CSV o Excel)
    cruzar_con_csv()          -> listas de control
    exportar_final()          -> archivo final: columnas del CSV + horarios (solo valores)

Utilidades independientes:
    util_corregir_nitro(), util_personal_en_firma() (macro 2), util_buscar_gua() (macro 8), util_formato_quincena() (macro 9)
"""
from __future__ import annotations

import copy
import csv
import io
import re
import unicodedata
from datetime import datetime
from datetime import time as dtime
from difflib import SequenceMatcher

import pandas as pd
import pypdf
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --------------------------------------------------------------------------
# Constantes
# --------------------------------------------------------------------------
DIAS = ["L", "M", "Mi", "J", "V", "S", "D"]
COLS_HORA = ["ENTRADA", "SALIDA", "HORAS DE TRABAJO"]

# Columnas C a Q del layout de horarios (15 columnas). TOLERANCIAS ocupa la N.
COLS_C_A_Q = [
    "ENTRADA",
    "SALIDA",
    *DIAS,
    "HORAS DE TRABAJO",
    "TURNO",
    "TOLERANCIAS",
    "DESCANSO",
    "TRABAJO",
    "TOTAL",
]
ORDEN_DATOS = ["TARJETA", "NOMBRE", *COLS_C_A_Q]

# Solo para la utilidad «Obtener personal EN FIRMA» (macro 2), que sí usa la quincena.
# (nombre en el reporte, encabezados aceptados en la quincena, opcional)
COL_QUINCENA = [
    ("RFC", ["rfc"], False),
    ("CURP", ["curp"], False),
    ("SEXO", ["sexo"], True),
    ("UR", ["ur"], False),
    ("NOMBRE DEL PROGRAMA", ["nombre del programa", "programa"], True),
    ("PUESTO", ["puesto"], False),
    ("DESCRIPCION", ["descripcion", "descripción"], True),
    (
        "CLAVE PRESUPUESTAL",
        ["cve_pago", "clave presupuestal", "clave_presupuestal", "cve pago"],
        False,
    ),
    ("CENTRO", ["centro"], False),
    ("CLUES", ["clues"], False),
    ("REGION", ["region", "región"], True),
    ("CONTRATO", ["contrato"], False),
    ("TIPO DE TRABAJADOR", ["tipo de trabajador", "tipo_trabajador"], True),
    ("RAMA", ["rama"], True),
]
COLS_EN_FIRMA = [
    "TARJETA",
    "RFC",
    "CURP",
    "NOMBRE",
    "UR",
    "PUESTO",
    "CLAVE PRESUPUESTAL",
    "CENTRO",
    "CLUES",
    "CONTRATO",
]

C_ENCABEZADO = "A3C2AD"  # RGB(163,194,173), el verde de encabezado de las macros


def _avanzar(cb, frac, texto=""):
  """Reporta avance (0 a 1) a la barra de progreso, si se pasó una."""
  if cb is not None:
    cb(min(max(frac, 0.0), 1.0), texto)


def _relleno(hex_color: str) -> PatternFill:
  return PatternFill(
      fill_type="solid", start_color=hex_color, end_color=hex_color
  )


# --------------------------------------------------------------------------
# Utilidades de lectura / normalización
# --------------------------------------------------------------------------
def es_vacio(v) -> bool:
  if v is None:
    return True
  try:
    if pd.isna(v):
      return True
  except (TypeError, ValueError):
    pass
  return isinstance(v, str) and v.strip() == ""


def celda(fila, i):
  return fila[i] if i < len(fila) else None


def valor(v):
  """Valor de celda limpio: texto sin espacios sobrantes, números tal cual."""
  if es_vacio(v):
    return None
  return v.strip() if isinstance(v, str) else v


def clave(v) -> str:
  """Normaliza una tarjeta para comparar (1023, 1023.0, '1023', '01023' -> '1023')."""
  if es_vacio(v):
    return ""
  if isinstance(v, float) and v.is_integer():
    v = int(v)
  s = str(v).strip()
  if s.isdigit():
    s = str(int(s))
  return s


def _unicos(etiquetas):
  """Evita nombres de columna repetidos (el CSV de MySQL trae dos «NOMBRE»)."""
  vistos, salida = {}, []
  for e in etiquetas:
    n = vistos.get(e, 0)
    vistos[e] = n + 1
    salida.append(e if n == 0 else f"{e}.{n}")
  return salida


def leer_hoja1(archivo) -> list[list]:
  """Devuelve todas las filas de la primera hoja del libro."""
  if hasattr(archivo, "seek"):
    archivo.seek(0)
  wb = load_workbook(archivo, read_only=True, data_only=True)
  try:
    return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
  finally:
    wb.close()


_RE_HORA = re.compile(r"^\s*(\d{1,2})[:\.](\d{2})(?::\d{2})?\s*$")


def hora_a_min(v):
  """Minutos desde las 00:00, o None si no es una hora válida."""
  if es_vacio(v):
    return None
  if isinstance(v, (datetime, dtime)):
    return v.hour * 60 + v.minute
  if isinstance(v, (int, float)) and not isinstance(v, bool):
    return int(round(v * 1440)) if 0 <= v < 1 else None  # fracción de día de Excel
  m = _RE_HORA.match(str(v))
  if m:
    h, mi = int(m.group(1)), int(m.group(2))
    if h < 24 and mi < 60:
      return h * 60 + mi
  return None


def normalizar_hora(v):
  """'7:00' -> '07:00'. Lo que no sea hora se conserva como texto (para que salga REVISAR)."""
  m = hora_a_min(v)
  if m is not None:
    return f"{m // 60:02d}:{m % 60:02d}"
  return None if es_vacio(v) else str(v).strip()


def _a_hora_excel(v):
  m = hora_a_min(v)
  if m is None:
    return None if es_vacio(v) else v
  return dtime(m // 60, m % 60)


def _autoajustar(ws, minimo=8, maximo=60):
  anchos = {}
  for fila in ws.iter_rows():
    for c in fila:
      if c.value is None:
        continue
      largo = 5 if isinstance(c.value, (dtime, datetime)) else len(str(c.value))
      anchos[c.column_letter] = max(anchos.get(c.column_letter, 0), largo)
  for col, w in anchos.items():
    ws.column_dimensions[col].width = min(max(w + 2, minimo), maximo)


# --------------------------------------------------------------------------
# Reglas de turno (macros 5 y 6)
# --------------------------------------------------------------------------
def _m(hhmm: str) -> int:
  h, mi = hhmm.split(":")
  return int(h) * 60 + int(mi)


def clasificar_turno(entrada, horas, dias) -> str:
  """Replica las reglas de las macros. `dias` = [L, M, Mi, J, V, S, D].

  c = hora de entrada, l = horas de trabajo (jornada). Rangos inclusivos.
  El orden de las reglas importa: se aplica la primera que cumple.
  """
  c, l = hora_a_min(entrada), hora_a_min(horas)
  if c is None or l is None:
    return "REVISAR"
  vacios = sum(es_vacio(d) for d in dias)  # CountBlank(E:K)
  e = dias[0]  # lunes
  j = None if es_vacio(dias[5]) else str(dias[5]).strip()  # sábado
  k = None if es_vacio(dias[6]) else str(dias[6]).strip()  # domingo

  if _m("05:00") <= c <= _m("11:30") and _m("05:30") <= l <= _m("09:00"):
    return "MATUTINO"
  if _m("12:00") <= c <= _m("17:30") and _m("05:30") <= l <= _m("09:00"):
    return "VESPERTINO"
  if (
      _m("18:00") <= c <= _m("22:00")
      and _m("10:00") <= l <= _m("12:00")
      and vacios > 1
  ):
    return "NOCTURNO"
  if _m("06:00") <= c <= _m("11:00") and _m("10:00") <= l <= _m("12:00"):
    return "JORNADA ACUMULADA"
  if (
      _m("18:00") <= c <= _m("22:00")
      and _m("10:00") <= l <= _m("12:00")
      and es_vacio(e)
      and j == "*"
      and k == "*"
  ):
    return "JORNADA ACUMULADA NOCTURNA"
  return "REVISAR"


def evaluar_turnos(df: pd.DataFrame) -> pd.DataFrame:
  """Normaliza horas/días y recalcula TURNO en todas las filas (macro 6)."""
  df = df.copy()
  for c in COLS_HORA:
    df[c] = df[c].map(normalizar_hora).astype(object)
  for d in DIAS:
    df[d] = (
        df[d].map(lambda v: None if es_vacio(v) else str(v).strip()).astype(object)
    )
  df["TURNO"] = df.apply(
      lambda r: clasificar_turno(
          r["ENTRADA"], r["HORAS DE TRABAJO"], [r[d] for d in DIAS]
      ),
      axis=1,
  ).astype(object)
  return df


# --------------------------------------------------------------------------
# Reporte de Inspector (bloques de 5 filas por empleado)
# --------------------------------------------------------------------------
_RE_EMPLEADO = re.compile(r"^\s*(\d+)\s+(\S.*?)\s*$")
_RE_TOTAL = re.compile(r"\bTotal\s*=", re.I)


def partir_total(texto: str):
  """'Descanso = 15h  Trabajo = 37h 30m  Total = 52:30' -> ('15h', '37h 30m', '52:30').

  Si falta alguna parte (p. ej. empleados sin descanso) devuelve '' en esa posición.
  """
  t = re.sub(r"\s+", " ", texto or "").strip()
  d = re.search(r"Descanso\s*=\s*(.*?)(?=\s+Trabajo\s*=|\s+Total\s*=|$)", t, re.I)
  w = re.search(r"Trabajo\s*=\s*(.*?)(?=\s+Total\s*=|$)", t, re.I)
  o = re.search(r"Total\s*=\s*(\S+)", t, re.I)
  return (
      d.group(1).strip() if d else "",
      w.group(1).strip() if w else "",
      o.group(1).strip() if o else "",
  )


def leer_reporte_inspector(archivo, etiqueta: str):
  """Lee un reporte de Inspector convertido a Excel (Hoja1).

  No depende de saltos fijos de 5 filas: reconoce cada empleado por su renglón
  «número  NOMBRE» y toma el renglón de horario que sigue al encabezado
  «Actividad».
  Ignora los encabezados «Reporte de horarios por empleado» que se repiten por
  página.
  TOLERANCIAS es el valor de la columna L de ese renglón de horario (L4, L9,
  L14… del layout).
  """
  filas = leer_hoja1(archivo)
  empleados, emp = [], None
  for r in filas:
    a, b = celda(r, 0), celda(r, 1)
    if isinstance(a, str):
      sa = a.strip()
      if sa.lower().startswith("reporte de horarios"):
        continue
      if _RE_TOTAL.search(sa):
        if emp is not None:
          emp["total"] = sa
        continue
      m = _RE_EMPLEADO.match(sa)
      if m:
        emp = {
            "tarjeta": clave(m.group(1)),
            "nombre": m.group(2),
            "horario": None,
            "hdr": False,
            "trabajos": [],
            "total": "",
        }
        empleados.append(emp)
        continue
    if emp is None or es_vacio(b):
      continue
    sb = str(b).strip()
    if sb.lower() == "actividad":
      emp["hdr"] = True
    elif emp["hdr"]:
      if (
          hora_a_min(celda(r, 2)) is not None
      ):  # renglón de horario (Trabajo, trabajo...)
        emp["trabajos"].append(r)
    elif emp["horario"] is None:
      emp["horario"] = sb  # texto descriptivo del horario

  avisos, registros, vistos = [], [], set()
  for e in empleados:
    t = e["trabajos"]
    if e["tarjeta"] in vistos:
      avisos.append(
          f"[{etiqueta}] Tarjeta {e['tarjeta']} ({e['nombre']}) aparece más de"
          " una vez; se usó la primera."
      )
      continue
    vistos.add(e["tarjeta"])
    if len(t) > 1:
      avisos.append(
          f"[{etiqueta}] Tarjeta {e['tarjeta']} ({e['nombre']}) tiene {len(t)}"
          " horarios; se usó el primero. Revísalo."
      )
    elif not t:
      avisos.append(
          f"[{etiqueta}] Tarjeta {e['tarjeta']} ({e['nombre']}) no trae renglón"
          " de horario."
      )
    r = t[0] if t else []
    desc, trab, tot = partir_total(e["total"])
    registros.append({
        "TARJETA": e["tarjeta"],
        "NOMBRE": e["nombre"],
        "ENTRADA": normalizar_hora(celda(r, 2)),
        "SALIDA": normalizar_hora(celda(r, 3)),
        **{
            d: (
                None
                if es_vacio(celda(r, 4 + i))
                else str(celda(r, 4 + i)).strip()
            )
            for i, d in enumerate(DIAS)
        },
        "HORAS DE TRABAJO": normalizar_hora(celda(r, 14)),
        "TOLERANCIAS": valor(celda(r, 11)),
        "DESCANSO": f"Descanso ={desc}" if desc else None,
        "TRABAJO": f"Trabajo ={trab}" if trab else None,
        "TOTAL": f"Total ={tot}" if tot else None,
        "_HORARIO": e["horario"],
    })
  if not registros:
    raise ValueError(
        f"No encontré empleados en el reporte de {etiqueta}. Revisa que sea el"
        " reporte de Inspector convertido a Excel (Hoja1)."
    )
  return registros, avisos


def construir_datos(act, fir):
  """Une activos y en firma en una sola tabla y calcula el turno."""
  avisos = []
  firma_set = {r["TARJETA"] for r in fir}
  registros, ya = list(act), {r["TARJETA"] for r in act}
  for r in fir:
    if r["TARJETA"] in ya:
      avisos.append(
          f"Tarjeta {r['TARJETA']} ({r['NOMBRE']}) está en ACTIVOS y EN FIRMA;"
          " se dejó una sola fila."
      )
    else:
      registros.append(r)
      ya.add(r["TARJETA"])
  filas = [{**r, "_INACTIVO": r["TARJETA"] in firma_set} for r in registros]
  df = pd.DataFrame(
      filas, columns=[*ORDEN_DATOS, "_HORARIO", "_INACTIVO"], dtype=object
  )
  return evaluar_turnos(df), avisos


# --------------------------------------------------------------------------
# Export de MySQL (CSV o Excel) y cruce con los horarios
# --------------------------------------------------------------------------
def _decodificar(raw: bytes) -> str:
  for enc in ("utf-8-sig", "cp1252"):
    try:
      return raw.decode(enc)
    except UnicodeDecodeError:
      continue
  return raw.decode("latin-1")


def leer_tabla(archivo) -> list[list]:
  """Filas de un .csv (separador y codificación detectados) o de la Hoja1 de un .xlsx."""
  nombre = str(getattr(archivo, "name", archivo)).lower()
  if not nombre.endswith((".csv", ".txt")):
    return leer_hoja1(archivo)
  if hasattr(archivo, "read"):
    archivo.seek(0)
    raw = archivo.read()
  else:
    with open(archivo, "rb") as f:
      raw = f.read()
  texto = _decodificar(raw)
  primera = next((ln for ln in texto.splitlines() if ln.strip(", ;\t|")), "")
  sep = max([",", ";", "\t", "|"], key=primera.count)
  return [
      list(r) for r in csv.reader(io.StringIO(texto, newline=""), delimiter=sep)
  ]


def _convertir(v):
  """Texto -> número cuando es un entero sin ceros a la izquierda (como lo abriría Excel)."""
  if es_vacio(v):
    return None
  if isinstance(v, str):
    s = v.strip()
    return int(s) if re.fullmatch(r"0|[1-9]\d{0,14}", s) else s
  return v


def leer_csv_mysql(archivo):
  """Lee el export de MySQL. La columna A debe ser la tarjeta; se ignoran las filas vacías."""
  filas = leer_tabla(archivo)
  i0 = next(
      (i for i, r in enumerate(filas) if any(not es_vacio(c) for c in r)), None
  )
  if i0 is None:
    raise ValueError("El archivo de MySQL está vacío.")
  enc = ["" if es_vacio(c) else str(c).strip() for c in filas[i0]]
  while enc and not enc[-1]:
    enc.pop()
  if not enc or "tarjeta" not in enc[0].lower():
    raise ValueError(
        "En el archivo de MySQL, la columna A debe ser la tarjeta"
        " (TARJETAEMPLEADO) y el encabezado debe estar en la primera fila."
    )
  n = len(enc)
  datos, vacias = [], 0
  for r in filas[i0 + 1 :]:
    r = list(r)[:n]
    if all(es_vacio(c) for c in r):
      vacias += 1
      continue
    datos.append([_convertir(c) for c in r] + [None] * (n - len(r)))
  return {"encabezado": enc, "filas": datos, "vacias": vacias}


def cruzar_con_csv(df: pd.DataFrame, csv_info: dict) -> dict:
  """Listas de control: filas del CSV sin coincidencia y empleados de Inspector fuera del CSV."""
  enc, filas = csv_info["encabezado"], csv_info["filas"]
  en_inspector = set(df["TARJETA"])
  claves_csv = [clave(f[0]) for f in filas]
  vistos, repetidas = set(), set()
  for k in claves_csv:
    if k and k in vistos:
      repetidas.add(k)
    vistos.add(k)

  idx_vis = [i for i in (0, 2, 3) if i < len(enc)]
  etiquetas = _unicos([enc[i] for i in idx_vis])
  sin = [
      [f[i] for i in idx_vis]
      for f, k in zip(filas, claves_csv)
      if k not in en_inspector
  ]
  coinc = sum(1 for k in claves_csv if k in en_inspector)
  no_en_csv = df.loc[~df["TARJETA"].isin(set(claves_csv)), ["TARJETA", "NOMBRE"]]

  avisos = []
  if repetidas:
    avisos.append(
        f"{len(repetidas)} tarjeta(s) repetida(s) en el CSV (cada fila recibe"
        " el mismo horario): "
        + ", ".join(sorted(repetidas)[:15])
        + ("…" if len(repetidas) > 15 else "")
    )
  j = next(
      (i for i, h in enumerate(enc) if h.strip().lower() == "borrado"), None
  )
  if j is not None:
    marcadas = sum(
        1
        for f, k in zip(filas, claves_csv)
        if k in en_inspector and not es_vacio(f[j]) and str(f[j]).strip() != "0"
    )
    if marcadas:
      avisos.append(
          f"{marcadas} fila(s) con coincidencia tienen BORRADO distinto de 0;"
          " se incluyen en el archivo final (puedes filtrarlas por esa"
          " columna)."
      )
  if len(enc) > 9:
    avisos.append(
        f"El CSV trae {len(enc)} columnas; los horarios se pegan desde la"
        f" columna {get_column_letter(len(enc) + 1)} para no sobrescribir"
        " datos."
    )
  return {
      "coincidencias": coinc,
      "sin_coincidencia": pd.DataFrame(sin, columns=etiquetas, dtype=object),
      "no_en_csv": no_en_csv.reset_index(drop=True),
      "vacias": csv_info["vacias"],
      "avisos": avisos,
  }


def procesar(f_activos, f_firma, f_csv, progreso=None):
  """`f_firma` es opcional (None): sin él solo se trabaja con los empleados del reporte de activos."""
  _avanzar(progreso, 0.05, "Leyendo reporte de activos…")
  act, av1 = leer_reporte_inspector(f_activos, "activos")
  _avanzar(progreso, 0.35, "Leyendo reporte en firma…")
  fir, av2 = (
      leer_reporte_inspector(f_firma, "en firma")
      if f_firma is not None
      else ([], [])
  )
  _avanzar(progreso, 0.60, "Calculando turnos…")
  df, av3 = construir_datos(act, fir)
  _avanzar(progreso, 0.75, "Leyendo CSV de MySQL…")
  csv_info = leer_csv_mysql(f_csv)
  _avanzar(progreso, 0.90, "Cruzando información…")
  cruce = cruzar_con_csv(df, csv_info)
  _avanzar(progreso, 1.0, "Listo")
  return {
      "datos": df,
      "csv": csv_info,
      "cruce": cruce,
      "avisos": av1 + av2 + av3 + cruce["avisos"],
      "version": 0,
      "final": None,
  }


# --------------------------------------------------------------------------
# Archivo final
# --------------------------------------------------------------------------
def exportar_final(df: pd.DataFrame, csv_info: dict) -> bytes:
  """Filas del CSV con coincidencia: columnas del CSV (A:I) + columnas C:Q del reporte (J:X).

  Solo valores. Si el CSV tuviera más de 9 columnas, los horarios se pegan después de la última.
  """
  enc, filas = csv_info["encabezado"], csv_info["filas"]
  n = len(enc)
  inicio = max(10, n + 1)  # columna (1-based) donde empiezan los horarios
  por_tarjeta = {
      r["TARJETA"]: [r[c] for c in COLS_C_A_Q] for _, r in df.iterrows()
  }
  hora_idx = {inicio + COLS_C_A_Q.index(c) for c in COLS_HORA}

  wb = Workbook()
  ws = wb.active
  ws.title = "REPORTE_HORARIOS"
  ws.append(enc + [None] * (inicio - 1 - n) + COLS_C_A_Q)
  for c in range(1, inicio + len(COLS_C_A_Q)):
    x = ws.cell(row=1, column=c)
    x.fill = _relleno(C_ENCABEZADO)
    x.font = Font(bold=True)
    x.alignment = Alignment(horizontal="center", vertical="center")

  for f in filas:
    h = por_tarjeta.get(clave(f[0]))
    if h is None:
      continue
    valores = [None if es_vacio(v) else v for v in h]
    for i, c in enumerate(COLS_C_A_Q):
      if c in COLS_HORA:
        valores[i] = _a_hora_excel(valores[i])
    ws.append(f + [None] * (inicio - 1 - n) + valores)
    for col in hora_idx:
      x = ws.cell(row=ws.max_row, column=col)
      if isinstance(x.value, dtime):
        x.number_format = "hh:mm"
  ws.freeze_panes = "A2"
  _autoajustar(ws)
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue()


# --------------------------------------------------------------------------
# Utilidades independientes
# --------------------------------------------------------------------------
def util_corregir_nitro(pdf_file, excel_file, progreso=None):
  """Corrige la omisión de la primera línea de empleado que Nitro realiza por cada página del PDF."""
  reader = pypdf.PdfReader(pdf_file)
  patron_emp = re.compile(r"^(\d{1,6})\s+([A-ZÑÁÉÍÓÚ\.\-\,\s]+)$")

  # 1. Extraer secuencia completa del PDF
  secuencia_pdf = []
  n_pag = len(reader.pages)
  for i, page in enumerate(reader.pages, start=1):
    text = page.extract_text() or ""
    for line in text.split("\n"):
      line_str = line.strip()
      if patron_emp.match(line_str):
        secuencia_pdf.append(line_str)
    if i % 5 == 0 or i == n_pag:
      _avanzar(progreso, 0.60 * i / n_pag, f"Leyendo PDF: página {i} de {n_pag}")

  # 2. Cargar Excel original de Nitro
  _avanzar(progreso, 0.62, "Abriendo Excel de Nitro…")
  if hasattr(excel_file, "seek"):
    excel_file.seek(0)
  wb_orig = load_workbook(excel_file)
  ws_orig = wb_orig.active

  # 3. Mapear posiciones faltantes
  insert_dict = {}
  idx_pdf = 0

  r = 1
  while r <= ws_orig.max_row and idx_pdf < len(secuencia_pdf):
    val_a = str(ws_orig.cell(r, 1).value or "").strip()
    val_m = str(ws_orig.cell(r, 13).value or "").strip()

    if patron_emp.match(val_a):
      idx_pdf += 1
      r += 1
    elif r == 1 or val_m == "Fecha":
      prev_a = str(ws_orig.cell(r - 1, 1).value or "").strip() if r > 1 else ""
      if not patron_emp.match(prev_a):
        insert_dict[r] = secuencia_pdf[idx_pdf]
        idx_pdf += 1
      r += 1
    else:
      r += 1

  # 4. Reconstruir nuevo libro manteniendo formatos e insertar filas faltantes
  wb_new = Workbook()
  ws_new = wb_new.active
  ws_new.title = ws_orig.title

  for col_letter, dim in ws_orig.column_dimensions.items():
    ws_new.column_dimensions[col_letter].width = dim.width

  new_row_idx = 1
  for orig_r in range(1, ws_orig.max_row + 1):
    if orig_r % 200 == 0:
      _avanzar(
          progreso,
          0.65 + 0.33 * orig_r / ws_orig.max_row,
          f"Reconstruyendo Excel: fila {orig_r} de {ws_orig.max_row}",
      )
    if orig_r in insert_dict:
      ws_new.cell(row=new_row_idx, column=1, value=insert_dict[orig_r])
      new_row_idx += 1

    for col_c in range(1, ws_orig.max_column + 1):
      src_cell = ws_orig.cell(row=orig_r, column=col_c)
      dest_cell = ws_new.cell(row=new_row_idx, column=col_c)

      dest_cell.value = src_cell.value
      if src_cell.has_style:
        dest_cell.font = copy.copy(src_cell.font)
        dest_cell.fill = copy.copy(src_cell.fill)
        dest_cell.border = copy.copy(src_cell.border)
        dest_cell.alignment = copy.copy(src_cell.alignment)
        dest_cell.number_format = src_cell.number_format
        dest_cell.protection = copy.copy(src_cell.protection)

    new_row_idx += 1

  # 5. Remaquetar celdas combinadas recalculando rangos
  for merged_range in list(ws_orig.merged_cells.ranges):
    min_col, min_row, max_col, max_row = merged_range.bounds

    shift_before = sum(1 for row_k in insert_dict.keys() if row_k <= min_row)
    shift_total = sum(1 for row_k in insert_dict.keys() if row_k <= max_row)

    new_min_row = min_row + shift_before
    new_max_row = max_row + shift_total

    new_range_str = f"{get_column_letter(min_col)}{new_min_row}:{get_column_letter(max_col)}{new_max_row}"
    ws_new.merge_cells(new_range_str)

  _avanzar(progreso, 0.99, "Guardando archivo…")
  buf = io.BytesIO()
  wb_new.save(buf)
  _avanzar(progreso, 1.0, "Listo")
  return buf.getvalue(), len(insert_dict)


# --------------------------------------------------------------------------
# Utilidad: agrupar empleados por tipo de horario (sobre el reporte final)
# --------------------------------------------------------------------------
_ORDEN_TIPOS = ["BASIFICADOS", "EVENTUALES", "GUARDERIA", "SIN TOLERANCIA"]
_ORDEN_CLASES = [
    "MATUTINO",
    "VESPERTINO",
    "NOCTURNO",
    "JORNADA ACUMULADA DIURNA",
    "JORNADA ACUMULADA NOCTURNA",
    "IRREGULAR",
    "SIN HORARIO",
]
_ORDEN_TURNOS = [
    "MATUTINO",
    "VESPERTINO",
    "NOCTURNO",
    "JORNADA ACUMULADA",
    "JORNADA ACUMULADA NOCTURNA",
    "REVISAR",
]
_DIAS_LARGO = [
    "LUNES", "MARTES", "MIERCOLES", "JUEVES", "VIERNES", "SABADO", "DOMINGO",
]
_DIAS_CORTO = ["L", "M", "Mi", "J", "V", "S", "D"]
# Los 7 patrones nocturnos válidos: 3 * (cada uno seguido de celda vacía, la
# salida) y una sola D. La letra la define el día que descansa.
_NOCTURNOS = {
    "*-*-*-D": "A",
    "D*-*-*-": "B",
    "-*-*-D*": "C",
    "*-*-D*-": "D",
    "*-D*-*-": "E",
    "-D*-*-*": "F",
    "-*-D*-*": "G",
}
C_SIN_TOLERANCIA = "F8CBAD"  # naranja suave: tolerancia vacía
C_TOLERANCIA_RARA = "FF6B6B"  # rojo: tolerancia distinta a las 3 principales
C_IRREGULAR = "FFE699"  # amarillo: días irregulares
C_GRUPO_ALTERNO = "EAF1F8"  # azul muy suave: sombreado alterno por grupo


def _tipo_tolerancia(v):
  """TOLERANCIAS -> palabra del distintivo (GENERALES SSJ = BASIFICADOS)."""
  if es_vacio(v):
    return "SIN TOLERANCIA"
  s = unicodedata.normalize("NFKD", str(v)).encode("ascii", "ignore").decode()
  s = s.strip().lower()
  if s.startswith("general"):
    return "BASIFICADOS"
  if s.startswith("eventual"):
    return "EVENTUALES"
  if s.startswith("guarder"):
    return "GUARDERIA"
  return str(v).strip().upper()  # otra tolerancia: se deja como viene (en rojo)


def _valor_comparable(v):
  """Valor normalizado para comparar celdas (las horas se comparan por minutos)."""
  if es_vacio(v):
    return None
  if isinstance(v, (datetime, dtime, str)):
    m = hora_a_min(v)
    if m is not None:
      return ("h", m)
  if isinstance(v, str):
    return v.strip()
  return v


def _minutos(v):
  return hora_a_min(v) if isinstance(v, (datetime, dtime, str)) else None


def _min_o_final(v):
  m = _minutos(v)
  return 9999 if m is None else m


def _patron_dias(valores):
  """Marcas L..D -> texto de 7 caracteres ('*', 'D', '-' = vacía) o None si hay otra cosa."""
  marcas = []
  for v in valores:
    if es_vacio(v):
      marcas.append("-")
    else:
      m = str(v).strip().upper()
      if m not in ("*", "D"):
        return None
      marcas.append(m)
  return "".join(marcas)


def _inicio_corrida(patron):
  """Jornada corrida = 5 * seguidos (en círculo) y 2 D seguidas. Devuelve el día de inicio."""
  if patron.count("*") != 5 or patron.count("D") != 2 or "-" in patron:
    return None
  for i in range(7):
    if patron[i - 1] == "D" and all(patron[(i + k) % 7] == "*" for k in range(5)):
      return i
  return None


def _describir(valores, entrada, turno_t):
  """Clasifica un horario por la estructura de sus días y la hora de entrada.

  Devuelve (clase, prefijo, dias, subclave). `clase` ordena los grupos; `prefijo` y
  `dias` forman la etiqueta. La columna TURNO solo se usa para rotular irregulares.
  """
  c = _minutos(entrada)
  if c is None:
    return "SIN HORARIO", "SIN HORARIO", "", ()
  diurna = 5 * 60 <= c < 18 * 60
  patron = _patron_dias(valores)
  if patron is not None:
    ini = _inicio_corrida(patron)
    if ini is not None and diurna:
      clase = "MATUTINO" if c < 12 * 60 else "VESPERTINO"
      return clase, clase, f"{_DIAS_LARGO[ini]} A {_DIAS_LARGO[(ini + 4) % 7]}", ini
    letra = _NOCTURNOS.get(patron)
    if letra is not None and not diurna:
      dias = "-".join(_DIAS_CORTO[i] for i, m in enumerate(patron) if m == "*")
      return "NOCTURNO", "NOCTURNO", f"{letra} ({dias})", letra
    if patron == "DDDDD**" and diurna:
      return "JORNADA ACUMULADA DIURNA", "JORNADA ACUMULADA DIURNA", "", ()
    if patron == "-DDDD**" and not diurna:
      return "JORNADA ACUMULADA NOCTURNA", "JORNADA ACUMULADA NOCTURNA", "", ()
  # Irregular: la estructura no corresponde a ninguna regla (o choca con la hora de entrada)
  dias = "-".join(
      _DIAS_CORTO[i] for i, v in enumerate(valores) if str(v or "").strip() == "*"
  ) or "SIN *"
  turno = "SIN TURNO" if es_vacio(turno_t) else str(turno_t).strip().upper()
  pos = _ORDEN_TURNOS.index(turno) if turno in _ORDEN_TURNOS else len(_ORDEN_TURNOS)
  return "IRREGULAR", f"IRREGULAR {turno}", dias, (pos, dias)


def util_agrupar_horarios(archivo):
  """Agrupa a los empleados del reporte final que tienen el mismo horario.

  Compara las columnas J a U (ENTRADA ... TOLERANCIAS). Copia las columnas A a X
  a una hoja nueva (AGRUPADO), con los empleados iguales juntos, y agrega en la
  columna Y el tipo de horario. Solo reordena y aplica formato de centrado y bordes.
  """
  if hasattr(archivo, "seek"):
    archivo.seek(0)
  wb = load_workbook(archivo)
  ws = wb.worksheets[0]
  if ws.max_row < 2:
    raise ValueError("El archivo no tiene filas de datos.")
  avisos = []

  # Columnas: se buscan por encabezado (ENTRADA ... TOLERANCIAS = J:U).
  nombres = [n.upper() for n in COLS_C_A_Q[:12]]
  ini = None
  for c in ws[1]:
    if str(c.value or "").strip().upper() == "ENTRADA":
      ini = c.column
      break
  if ini is None or [
      str(ws.cell(1, ini + i).value or "").strip().upper() for i in range(12)
  ] != nombres:
    ini = 10
    avisos.append(
        "No encontré los encabezados ENTRADA a TOLERANCIAS en la fila 1; se"
        " compararon las columnas J a U."
    )
  c_fin = ini + 14  # TOTAL (columna X)
  c_tipo = c_fin + 1  # distintivo (columna Y)
  c_dias = range(ini + 2, ini + 9)  # L..R
  c_turno, c_tol = ini + 10, ini + 11

  filas, vacias = [], 0
  for r in range(2, ws.max_row + 1):
    if all(es_vacio(ws.cell(r, c).value) for c in range(1, c_fin + 1)):
      vacias += 1
    else:
      filas.append(r)
  if not filas:
    raise ValueError("El archivo no tiene filas de datos.")
  if vacias:
    avisos.append(f"Se omitieron {vacias} fila(s) completamente vacías.")

  # Grupos: empleados idénticos en J:U
  grupos = {}
  for pos, r in enumerate(filas):
    k = tuple(_valor_comparable(ws.cell(r, c).value) for c in range(ini, ini + 12))
    g = grupos.get(k)
    if g is None:
      valores = [ws.cell(r, c).value for c in c_dias]
      clase, prefijo, dias, sub = _describir(
          valores, ws.cell(r, ini).value, ws.cell(r, c_turno).value
      )
      tipo = _tipo_tolerancia(ws.cell(r, c_tol).value)
      base = prefijo if clase == "SIN HORARIO" else " ".join(
          x for x in (prefijo, dias, tipo) if x
      )
      g = grupos[k] = {
          "clase": clase, "tipo": tipo, "sub": sub, "base": base,
          "filas": [], "primero": pos, "r0": r,
      }
    g["filas"].append(r)

  # Orden de los grupos y numeración de variantes
  por_base = {}
  for g in grupos.values():
    por_base.setdefault(g["base"], []).append(g)
  total_base = {b: sum(len(g["filas"]) for g in lista) for b, lista in por_base.items()}

  def _pos(lista, valor):
    return lista.index(valor) if valor in lista else len(lista)

  def _orden_base(item):
    base, lista = item
    g0 = lista[0]
    sub = g0["sub"]
    if g0["clase"] in ("MATUTINO", "VESPERTINO"):
      sub = (-total_base[base], sub)
    elif not isinstance(sub, tuple):
      sub = (sub,)
    return (
        _pos(_ORDEN_CLASES, g0["clase"]),
        _pos(_ORDEN_TIPOS, g0["tipo"]),
        g0["tipo"],
        sub,
        base,
    )

  ordenados = []
  for base, lista in sorted(por_base.items(), key=_orden_base):
    lista.sort(
        key=lambda g: (
            -len(g["filas"]),
            _min_o_final(ws.cell(g["r0"], ini).value),
            _min_o_final(ws.cell(g["r0"], ini + 1).value),
            g["primero"],
        )
    )
    for n, g in enumerate(lista, start=1):
      sin_numero = g["clase"] == "SIN HORARIO" and len(lista) == 1
      g["etiqueta"] = base if sin_numero else f"{base} {n}"
      ordenados.append(g)

  # Libro nuevo con una sola hoja
  wn = Workbook()
  wsn = wn.active
  wsn.title = "AGRUPADO"

  def _copiar(src, dst):
    dst.value = src.value
    if src.has_style:
      dst.font = copy.copy(src.font)
      dst.fill = copy.copy(src.fill)
      dst.number_format = copy.copy(src.number_format)
      dst.protection = copy.copy(src.protection)

  for c in range(1, c_fin + 1):
    _copiar(ws.cell(1, c), wsn.cell(1, c))
  enc_tipo = wsn.cell(1, c_tipo)
  _copiar(ws.cell(1, c_fin), enc_tipo)
  enc_tipo.value = "TIPO DE HORARIO"

  fila_nueva = 2
  n_irregulares = n_tol_rara = n_sin_tol = 0
  for i, g in enumerate(ordenados):
    irregular = g["clase"] == "IRREGULAR"
    sin_tol = g["tipo"] == "SIN TOLERANCIA" and g["clase"] != "SIN HORARIO"
    rara = g["tipo"] not in _ORDEN_TIPOS
    if irregular:
      n_irregulares += len(g["filas"])
    if rara:
      n_tol_rara += len(g["filas"])
    if sin_tol:
      n_sin_tol += len(g["filas"])
    for r in g["filas"]:
      for c in range(1, c_fin + 1):
        _copiar(ws.cell(r, c), wsn.cell(fila_nueva, c))
      wsn.cell(fila_nueva, c_tipo).value = g["etiqueta"]
      if i % 2 == 0:  # sombreado alterno por grupo
        for c in range(1, c_tipo + 1):
          cel = wsn.cell(fila_nueva, c)
          if cel.fill is None or cel.fill.fill_type is None:
            cel.fill = _relleno(C_GRUPO_ALTERNO)
      if irregular:  # días irregulares: amarillo
        for c in list(c_dias) + [c_tipo]:
          wsn.cell(fila_nueva, c).fill = _relleno(C_IRREGULAR)
      if sin_tol:  # tolerancia vacía: naranja
        for c in (c_tol, c_tipo):
          wsn.cell(fila_nueva, c).fill = _relleno(C_SIN_TOLERANCIA)
      if rara:  # tolerancia distinta a las 3 principales: rojo
        for c in (c_tol, c_tipo):
          wsn.cell(fila_nueva, c).fill = _relleno(C_TOLERANCIA_RARA)
      fila_nueva += 1

  # Aplicar formato unificado de centrado y bordes a toda la tabla
  thin_side = Side(style="thin", color="BFBFBF")
  table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
  center_align = Alignment(horizontal="center", vertical="center")

  for r_idx in range(1, fila_nueva):
    for c_idx in range(1, c_tipo + 1):
      cell = wsn.cell(row=r_idx, column=c_idx)
      cell.alignment = center_align
      cell.border = table_border

  for c in range(1, c_fin + 1):
    letra = get_column_letter(c)
    dim = ws.column_dimensions.get(letra)
    if dim is not None and dim.width:
      wsn.column_dimensions[letra].width = dim.width
  wsn.column_dimensions[get_column_letter(c_tipo)].width = (
      max(len(g["etiqueta"]) for g in ordenados) + 2
  )
  wsn.freeze_panes = "A2"

  # Verificación de integridad
  def _fila(hoja, r):
    return repr(tuple(_valor_comparable(hoja.cell(r, c).value) for c in range(1, c_fin + 1)))

  if sorted(_fila(ws, r) for r in filas) != sorted(
      _fila(wsn, r) for r in range(2, fila_nueva)
  ):
    raise RuntimeError(
        "La verificación falló: las filas del resultado no coinciden con el"
        " original. No se generó el archivo."
    )

  # Resto de avisos y retorno...
  if n_irregulares:
    avisos.append(
        f"{n_irregulares} empleado(s) con días irregulares (en amarillo): su"
        " combinación de * y D no corresponde a ninguna regla o no coincide con"
        " su hora de entrada."
    )
  if n_tol_rara:
    avisos.append(
        f"{n_tol_rara} empleado(s) con una tolerancia distinta a BASIFICADOS,"
        " EVENTUALES y GUARDERIA (en rojo)."
    )
  if n_sin_tol:
    avisos.append(f"{n_sin_tol} empleado(s) sin tolerancia (en naranja).")

  def _dias(r):
    return "".join(
        "-" if es_vacio(ws.cell(r, c).value) else str(ws.cell(r, c).value).strip()
        for c in c_dias
    )

  def _alerta(g):
    if g["clase"] == "IRREGULAR":
      return "DÍAS IRREGULARES"
    if g["tipo"] not in _ORDEN_TIPOS:
      return "TOLERANCIA DISTINTA"
    if g["tipo"] == "SIN TOLERANCIA" and g["clase"] != "SIN HORARIO":
      return "SIN TOLERANCIA"
    return ""

  resumen = pd.DataFrame([
      {
          "TIPO DE HORARIO": g["etiqueta"],
          "EMPLEADOS": len(g["filas"]),
          "ENTRADA": normalizar_hora(ws.cell(g["r0"], ini).value),
          "SALIDA": normalizar_hora(ws.cell(g["r0"], ini + 1).value),
          "DÍAS (L a D)": _dias(g["r0"]),
          "HORAS DE TRABAJO": normalizar_hora(ws.cell(g["r0"], ini + 9).value),
          "TOLERANCIAS": valor(ws.cell(g["r0"], c_tol).value),
          "ALERTA": _alerta(g),
      }
      for g in ordenados
  ])
  buf = io.BytesIO()
  wn.save(buf)
  return buf.getvalue(), resumen, avisos, len(filas)

# --------------------------------------------------------------------------
# Utilidad: agrupar por tipo de horario (columna Y) y UR (columna I)
# --------------------------------------------------------------------------
_CODIGOS_UR = ("FOR", "FO2", "FO3")


def _calcular_ur(v) -> str:
  """UR a partir de la columna G.

  Si la celda contiene FOR, FO2 o FO3 (en cualquier parte), la UR es solo ese
  código. En cualquier otro caso, son los 3 primeros caracteres.
  """
  if es_vacio(v):
    return ""
  if isinstance(v, float) and v.is_integer():
    v = int(v)
  texto = str(v).strip()
  hallados = [(texto.find(c), c) for c in _CODIGOS_UR if c in texto]
  if hallados:
    return min(hallados)[1]
  return texto[:3]


def _es_relleno_alterno(fill) -> bool:
  """True si el relleno es el azul suave del sombreado alterno por grupo."""
  try:
    if fill is None or fill.fill_type != "solid":
      return False
    rgb = fill.start_color.rgb
    return isinstance(rgb, str) and rgb.upper().endswith(C_GRUPO_ALTERNO)
  except Exception:  # noqa: BLE001
    return False


def util_agrupar_por_ur(archivo, progreso=None):
  """Escribe la UR en la columna I y agrupa por tipo de horario (Y) y UR (I).

  - Encabezado de I1 = "UR" (sobrescribe el existente) y cada celda de I se
    sobrescribe con la UR calculada desde la columna G.
  - Las filas se agrupan por TIPO DE HORARIO (columna Y, en el orden en que ya
    aparecen) y, dentro de cada tipo, por UR. Un grupo = misma Y y misma UR.
  - Resultado: libro nuevo con una hoja (AGRUPADO UR), columnas A a Y.
  Devuelve (bytes, resumen DataFrame, avisos, n_filas).
  """
  if hasattr(archivo, "seek"):
    archivo.seek(0)
  _avanzar(progreso, 0.05, "Abriendo archivo…")
  wb = load_workbook(archivo)
  ws = wb.worksheets[0]
  if ws.max_row < 2:
    raise ValueError("El archivo no tiene filas de datos.")
  avisos = []
  C_G, C_I = 7, 9

  # Columna Y: TIPO DE HORARIO
  c_tipo = None
  for c in ws[1]:
    if str(c.value or "").strip().upper() == "TIPO DE HORARIO":
      c_tipo = c.column
      break
  if c_tipo is None:
    if any(not es_vacio(ws.cell(r, 25).value) for r in range(2, ws.max_row + 1)):
      c_tipo = 25
      avisos.append(
          "No encontré el encabezado «TIPO DE HORARIO»; se usó la columna Y."
      )
    else:
      raise ValueError(
          "La columna Y (TIPO DE HORARIO) está vacía o no existe. Primero"
          " agrupa el reporte con «Agrupar empleados por tipo de horario»."
      )

  filas, vacias = [], 0
  for r in range(2, ws.max_row + 1):
    if all(es_vacio(ws.cell(r, c).value) for c in range(1, c_tipo + 1)):
      vacias += 1
    else:
      filas.append(r)
  if not filas:
    raise ValueError("El archivo no tiene filas de datos.")
  if vacias:
    avisos.append(f"Se omitieron {vacias} fila(s) completamente vacías.")

  # UR de cada fila y orden de los tipos de horario (el de aparición)
  _avanzar(progreso, 0.15, "Calculando UR…")
  info = {}
  orden_tipos = []
  sin_tipo = sin_ur = 0
  cuenta_codigos = {c: 0 for c in _CODIGOS_UR}
  for r in filas:
    tipo = valor(ws.cell(r, c_tipo).value)
    tipo = "" if es_vacio(tipo) else str(tipo).strip()
    ur = _calcular_ur(ws.cell(r, C_G).value)
    info[r] = (tipo, ur)
    if tipo and tipo not in orden_tipos:
      orden_tipos.append(tipo)
    if not tipo:
      sin_tipo += 1
    if not ur:
      sin_ur += 1
    if ur in cuenta_codigos:
      cuenta_codigos[ur] += 1
  if sin_tipo:
    avisos.append(
        f"{sin_tipo} fila(s) sin tipo de horario en la columna Y; quedaron al"
        " final."
    )
  if sin_ur:
    avisos.append(
        f"{sin_ur} fila(s) con la columna G vacía: su UR quedó vacía y van al"
        " final de su tipo de horario."
    )
  usados = [f"{n} con {c}" for c, n in cuenta_codigos.items() if n]
  if usados:
    avisos.append("UR asignada por código contenido en G: " + ", ".join(usados) + ".")

  def _clave(r):
    tipo, ur = info[r]
    pos = orden_tipos.index(tipo) if tipo else len(orden_tipos)
    return (pos, ur == "", ur, r)

  ordenadas = sorted(filas, key=_clave)

  # Libro nuevo
  wn = Workbook()
  wsn = wn.active
  wsn.title = "AGRUPADO UR"

  def _copiar(src, dst):
    dst.value = src.value
    if src.has_style:
      dst.font = copy.copy(src.font)
      if not _es_relleno_alterno(src.fill):
        dst.fill = copy.copy(src.fill)
      dst.number_format = copy.copy(src.number_format)
      dst.protection = copy.copy(src.protection)

  for c in range(1, c_tipo + 1):
    _copiar(ws.cell(1, c), wsn.cell(1, c))
  wsn.cell(1, C_I).value = "UR"

  fila_nueva = 2
  grupos = []
  grupo_idx = -1
  previo = None
  n_total = len(ordenadas)
  for i, r in enumerate(ordenadas, start=1):
    if i % 200 == 0:
      _avanzar(
          progreso, 0.20 + 0.70 * i / n_total, f"Copiando fila {i} de {n_total}"
      )
    tipo, ur = info[r]
    if (tipo, ur) != previo:
      grupo_idx += 1
      grupos.append({"tipo": tipo, "ur": ur, "n": 0})
      previo = (tipo, ur)
    grupos[-1]["n"] += 1
    for c in range(1, c_tipo + 1):
      _copiar(ws.cell(r, c), wsn.cell(fila_nueva, c))
    celda_ur = wsn.cell(fila_nueva, C_I)
    celda_ur.value = ur if ur else None
    celda_ur.number_format = "General"
    if grupo_idx % 2 == 0:  # sombreado alterno por grupo
      for c in range(1, c_tipo + 1):
        cel = wsn.cell(fila_nueva, c)
        if cel.fill is None or cel.fill.fill_type is None:
          cel.fill = _relleno(C_GRUPO_ALTERNO)
    fila_nueva += 1

  # Formato unificado de centrado y bordes
  thin_side = Side(style="thin", color="BFBFBF")
  table_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
  center_align = Alignment(horizontal="center", vertical="center")
  for r_idx in range(1, fila_nueva):
    for c_idx in range(1, c_tipo + 1):
      cell = wsn.cell(row=r_idx, column=c_idx)
      cell.alignment = center_align
      cell.border = table_border

  for c in range(1, c_tipo + 1):
    letra = get_column_letter(c)
    dim = ws.column_dimensions.get(letra)
    if dim is not None and dim.width:
      wsn.column_dimensions[letra].width = dim.width
  wsn.column_dimensions[get_column_letter(C_I)].width = max(
      wsn.column_dimensions[get_column_letter(C_I)].width or 0, 8
  )
  wsn.freeze_panes = "A2"

  # Verificación: mismas filas y valores, salvo la columna I
  _avanzar(progreso, 0.93, "Verificando integridad…")

  def _fila(hoja, r):
    return repr(
        tuple(
            _valor_comparable(hoja.cell(r, c).value)
            for c in range(1, c_tipo + 1)
            if c != C_I
        )
    )

  if sorted(_fila(ws, r) for r in filas) != sorted(
      _fila(wsn, r) for r in range(2, fila_nueva)
  ):
    raise RuntimeError(
        "La verificación falló: las filas del resultado no coinciden con el"
        " original. No se generó el archivo."
    )

  resumen = pd.DataFrame([
      {"TIPO DE HORARIO": g["tipo"], "UR": g["ur"], "EMPLEADOS": g["n"]}
      for g in grupos
  ])
  buf = io.BytesIO()
  wn.save(buf)
  _avanzar(progreso, 1.0, "Listo")
  return buf.getvalue(), resumen, avisos, len(filas)


# --------------------------------------------------------------------------
# Utilidad: agrupar por tipo GENERAL de horario + UR en pestañas separadas
# --------------------------------------------------------------------------
_ORDEN_CLASES_PEST = [
    "MATUTINO",
    "VESPERTINO",
    "NOCTURNO",
    "JORNADA ACUMULADA DIURNA",
    "JORNADA ACUMULADA NOCTURNA",
]
_TIPOS_CONOCIDOS = ("BASIFICADOS", "EVENTUALES", "GUARDERIA", "SIN TOLERANCIA")


def _tipo_general_y_orden(etiqueta: str):
  """Convierte la etiqueta detallada en un tipo general + clave de orden.

  - Quita el número final (1, 2, 3…).
  - Quita BASIFICADOS / EVENTUALES / GUARDERIA / SIN TOLERANCIA del final.
  - IRREGULAR, SIN HORARIO o vacío → («REVISAR», 99).
  """
  s = str(etiqueta or "").strip()
  su = s.upper()
  if (
      not s
      or su.startswith("IRREGULAR")
      or su == "SIN HORARIO"
      or su.startswith("SIN HORARIO")
  ):
    return "REVISAR", 99
  # Quitar número final
  s = re.sub(r"\s+\d+$", "", s).strip()
  # Quitar tipo conocido al final
  su = s.upper()
  for t in _TIPOS_CONOCIDOS:
    if su.endswith(" " + t):
      s = s[: -len(t) - 1].strip()
      break
  if not s:
    return "REVISAR", 99
  # Clave de orden según clase
  orden = 50
  su = s.upper()
  for i, pref in enumerate(_ORDEN_CLASES_PEST):
    if su.startswith(pref):
      orden = i
      break
  return s, orden


def _nombre_hoja_seguro(idx: int, tipo: str, ur: str) -> str:
  """Nombre de pestaña ≤ 31 caracteres y sin caracteres prohibidos por Excel."""
  if tipo == "REVISAR":
    return "REVISAR"
  t_corto = (tipo or "")[:18].strip()
  nombre = f"{idx:02d} {t_corto}"
  if ur:
    nombre = f"{nombre} | {ur}"
  nombre = nombre[:31]
  for c in r"[]:*?/\\":
    nombre = nombre.replace(c, "-")
  return nombre or f"Hoja{idx:02d}"


def util_agrupar_por_ur_pestanas(archivo, progreso=None):
  """Agrupa por tipo GENERAL de horario + UR y genera un libro con una pestaña
  por cada combinación (más una hoja ÍNDICE).

  - Unifica variantes numeradas y de tolerancia (BASIFICADOS/EVENTUALES/…).
    Ej.: «MATUTINO LUNES A VIERNES BASIFICADOS 1» y
    «MATUTINO LUNES A VIERNES EVENTUALES 3» → «MATUTINO LUNES A VIERNES».
  - Todo IRREGULAR, SIN HORARIO o vacío va a una sola pestaña «REVISAR».
  - Orden de pestañas: MATUTINO → VESPERTINO → NOCTURNO →
    JORNADA ACUMULADA DIURNA → JORNADA ACUMULADA NOCTURNA → REVISAR.
    Dentro de cada tipo general, las UR se ordenan alfabéticamente.
  - Columna I = «UR» (calculada igual que en util_agrupar_por_ur).
  - Devuelve (bytes, resumen DataFrame, avisos, n_filas).
  """
  if hasattr(archivo, "seek"):
    archivo.seek(0)
  _avanzar(progreso, 0.05, "Abriendo archivo…")
  wb = load_workbook(archivo)
  ws = wb.worksheets[0]
  if ws.max_row < 2:
    raise ValueError("El archivo no tiene filas de datos.")
  avisos = []
  C_G, C_I = 7, 9

  # Columna Y: TIPO DE HORARIO
  c_tipo = None
  for c in ws[1]:
    if str(c.value or "").strip().upper() == "TIPO DE HORARIO":
      c_tipo = c.column
      break
  if c_tipo is None:
    if any(not es_vacio(ws.cell(r, 25).value) for r in range(2, ws.max_row + 1)):
      c_tipo = 25
      avisos.append(
          "No encontré el encabezado «TIPO DE HORARIO»; se usó la columna Y."
      )
    else:
      raise ValueError(
          "La columna Y (TIPO DE HORARIO) está vacía o no existe. Primero"
          " agrupa el reporte con «Agrupar empleados por tipo de horario»."
      )

  filas, vacias = [], 0
  for r in range(2, ws.max_row + 1):
    if all(es_vacio(ws.cell(r, c).value) for c in range(1, c_tipo + 1)):
      vacias += 1
    else:
      filas.append(r)
  if not filas:
    raise ValueError("El archivo no tiene filas de datos.")
  if vacias:
    avisos.append(f"Se omitieron {vacias} fila(s) completamente vacías.")

  # Calcular tipo general + UR de cada fila
  _avanzar(progreso, 0.15, "Calculando tipos generales y UR…")
  info = {}
  sin_tipo = sin_ur = n_revisar = 0
  cuenta_codigos = {c: 0 for c in _CODIGOS_UR}
  for r in filas:
    tipo_orig = valor(ws.cell(r, c_tipo).value)
    tipo_orig = "" if es_vacio(tipo_orig) else str(tipo_orig).strip()
    tipo_gen, orden = _tipo_general_y_orden(tipo_orig)
    ur = _calcular_ur(ws.cell(r, C_G).value)
    # REVISAR no se separa por UR
    ur_grupo = "" if tipo_gen == "REVISAR" else ur
    info[r] = {
        "tipo_gen": tipo_gen,
        "orden": orden,
        "ur": ur,
        "ur_grupo": ur_grupo,
        "tipo_orig": tipo_orig,
    }
    if tipo_gen == "REVISAR":
      n_revisar += 1
    if not tipo_orig:
      sin_tipo += 1
    if not ur:
      sin_ur += 1
    if ur in cuenta_codigos:
      cuenta_codigos[ur] += 1

  if sin_tipo:
    avisos.append(
        f"{sin_tipo} fila(s) sin tipo de horario en la columna Y; se enviaron"
        " a la pestaña REVISAR."
    )
  if sin_ur:
    avisos.append(
        f"{sin_ur} fila(s) con la columna G vacía: su UR quedó vacía."
    )
  if n_revisar:
    avisos.append(
        f"{n_revisar} empleado(s) con horario irregular, sin horario o vacío"
        " se agruparon en la pestaña «REVISAR»."
    )
  usados = [f"{n} con {c}" for c, n in cuenta_codigos.items() if n]
  if usados:
    avisos.append(
        "UR asignada por código contenido en G: " + ", ".join(usados) + "."
    )

  # Ordenar filas: por orden de clase → tipo general → UR → fila original
  def _clave(r):
    i = info[r]
    return (i["orden"], i["tipo_gen"], i["ur_grupo"] == "", i["ur_grupo"], r)

  ordenadas = sorted(filas, key=_clave)

  # Construir grupos en el orden de aparición tras el sort
  grupos = []
  previo = None
  for r in ordenadas:
    i = info[r]
    clave_g = (i["tipo_gen"], i["ur_grupo"])
    if clave_g != previo:
      grupos.append({
          "tipo_gen": i["tipo_gen"],
          "ur": i["ur_grupo"],
          "filas": [],
          "n": 0,
      })
      previo = clave_g
    grupos[-1]["filas"].append(r)
    grupos[-1]["n"] += 1

  # Libro nuevo
  _avanzar(progreso, 0.25, "Creando pestañas…")
  wn = Workbook()
  # Hoja ÍNDICE
  ws_idx = wn.active
  ws_idx.title = "INDICE"

  # Encabezados del índice
  headers_idx = ["Nº", "TIPO GENERAL", "UR", "EMPLEADOS", "PESTAÑA"]
  for col, h in enumerate(headers_idx, start=1):
    cell = ws_idx.cell(1, col, value=h)
    cell.font = Font(bold=True)
    cell.fill = _relleno(C_ENCABEZADO)
    cell.alignment = Alignment(horizontal="center", vertical="center")

  thin_side = Side(style="thin", color="BFBFBF")
  table_border = Border(
      left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
  )
  center_align = Alignment(horizontal="center", vertical="center")

  def _copiar(src, dst):
    dst.value = src.value
    if src.has_style:
      dst.font = copy.copy(src.font)
      if not _es_relleno_alterno(src.fill):
        dst.fill = copy.copy(src.fill)
      dst.number_format = copy.copy(src.number_format)
      dst.protection = copy.copy(src.protection)

  n_total_filas = len(ordenadas)
  filas_copiadas = 0

  for g_idx, g in enumerate(grupos, start=1):
    nombre_hoja = _nombre_hoja_seguro(g_idx, g["tipo_gen"], g["ur"])
    # Evitar nombres duplicados (muy raro, pero por seguridad)
    base_nombre = nombre_hoja
    suf = 1
    while nombre_hoja in [s.title for s in wn.worksheets]:
      suf += 1
      nombre_hoja = (base_nombre[:28] + f"_{suf}")[:31]

    wsn = wn.create_sheet(title=nombre_hoja)

    # Encabezado
    for c in range(1, c_tipo + 1):
      _copiar(ws.cell(1, c), wsn.cell(1, c))
    wsn.cell(1, C_I).value = "UR"

    fila_nueva = 2
    for r in g["filas"]:
      filas_copiadas += 1
      if filas_copiadas % 150 == 0:
        _avanzar(
            progreso,
            0.25 + 0.65 * filas_copiadas / n_total_filas,
            f"Copiando fila {filas_copiadas} de {n_total_filas}",
        )
      for c in range(1, c_tipo + 1):
        _copiar(ws.cell(r, c), wsn.cell(fila_nueva, c))
      celda_ur = wsn.cell(fila_nueva, C_I)
      celda_ur.value = info[r]["ur"] if info[r]["ur"] else None
      celda_ur.number_format = "General"
      # Sombreado alterno por grupo (aquí cada pestaña es un grupo)
      if g_idx % 2 == 0:
        for c in range(1, c_tipo + 1):
          cel = wsn.cell(fila_nueva, c)
          if cel.fill is None or cel.fill.fill_type is None:
            cel.fill = _relleno(C_GRUPO_ALTERNO)
      fila_nueva += 1

    # Formato de la pestaña
    for r_idx in range(1, fila_nueva):
      for c_idx in range(1, c_tipo + 1):
        cell = wsn.cell(row=r_idx, column=c_idx)
        cell.alignment = center_align
        cell.border = table_border

    for c in range(1, c_tipo + 1):
      letra = get_column_letter(c)
      dim = ws.column_dimensions.get(letra)
      if dim is not None and dim.width:
        wsn.column_dimensions[letra].width = dim.width
    wsn.column_dimensions[get_column_letter(C_I)].width = max(
        wsn.column_dimensions[get_column_letter(C_I)].width or 0, 8
    )
    wsn.freeze_panes = "A2"

    # Fila en el índice
    row_idx = g_idx + 1
    ws_idx.cell(row_idx, 1, value=g_idx)
    ws_idx.cell(row_idx, 2, value=g["tipo_gen"])
    ws_idx.cell(row_idx, 3, value=g["ur"] if g["ur"] else None)
    ws_idx.cell(row_idx, 4, value=g["n"])
    link_cell = ws_idx.cell(row_idx, 5, value=f"→ {nombre_hoja}")
    link_cell.hyperlink = f"#'{nombre_hoja}'!A1"
    link_cell.font = Font(color="0563C1", underline="single")

  # Formato del índice
  for r_idx in range(1, len(grupos) + 2):
    for c_idx in range(1, 6):
      cell = ws_idx.cell(row=r_idx, column=c_idx)
      cell.alignment = center_align
      cell.border = table_border
  ws_idx.column_dimensions["A"].width = 6
  ws_idx.column_dimensions["B"].width = 45
  ws_idx.column_dimensions["C"].width = 12
  ws_idx.column_dimensions["D"].width = 12
  ws_idx.column_dimensions["E"].width = 28
  ws_idx.freeze_panes = "A2"

  # Verificación de integridad (mismas filas, salvo columna I)
  _avanzar(progreso, 0.95, "Verificando integridad…")

  def _fila(hoja, r):
    return repr(
        tuple(
            _valor_comparable(hoja.cell(r, c).value)
            for c in range(1, c_tipo + 1)
            if c != C_I
        )
    )

  filas_resultado = []
  for sheet in wn.worksheets:
    if sheet.title == "INDICE":
      continue
    for r in range(2, sheet.max_row + 1):
      if not all(
          es_vacio(sheet.cell(r, c).value) for c in range(1, c_tipo + 1)
      ):
        filas_resultado.append(_fila(sheet, r))

  if sorted(_fila(ws, r) for r in filas) != sorted(filas_resultado):
    raise RuntimeError(
        "La verificación falló: las filas del resultado no coinciden con el"
        " original. No se generó el archivo."
    )

  resumen = pd.DataFrame([
      {
          "TIPO GENERAL": g["tipo_gen"],
          "UR": g["ur"] if g["ur"] else "",
          "EMPLEADOS": g["n"],
          "PESTAÑA": _nombre_hoja_seguro(
              i + 1, g["tipo_gen"], g["ur"]
          ),
      }
      for i, g in enumerate(grupos)
  ])
  buf = io.BytesIO()
  wn.save(buf)
  _avanzar(progreso, 1.0, "Listo")
  return buf.getvalue(), resumen, avisos, len(filas)


def util_filtrar_por_quincenas(f_q1, f_q2, f_final):

  """Filtra el archivo final conservando formatos y estilos originales,
  comparando la columna B contra la columna A de la Quincena 1 y opcionalmente Q2,
  y copia las columnas PUESTO, DESCR, TIPOTRAB, RAMA y REGION a partir de la columna Z.
  """
  def _cargar_datos_quincena(archivo):
    if archivo is None:
      return {}
    if hasattr(archivo, "seek"):
      archivo.seek(0)
    wb_q = load_workbook(archivo, data_only=True)
    ws_q = wb_q.active
    
    idx_enc = None
    cols_map = {}
    target_headers = ["PUESTO", "DESCR", "TIPOTRAB", "RAMA", "REGION"]
    
    # Buscar dinámicamente la fila de encabezados y las columnas requeridas
    for r in range(1, min(15, ws_q.max_row + 1)):
      row_vals = [ws_q.cell(r, c).value for c in range(1, ws_q.max_column + 1)]
      row_strs = [str(v).strip().upper() if v is not None else "" for v in row_vals]
      if "TARJETA" in row_strs:
        idx_enc = r
        for idx_c, val in enumerate(row_strs, start=1):
          if val == "TARJETA":
            cols_map["TARJETA"] = idx_c
          elif val in target_headers:
            cols_map[val] = idx_c
        break
        
    if idx_enc is None or "TARJETA" not in cols_map:
      idx_enc = 1
      cols_map["TARJETA"] = 1
      for c in range(1, ws_q.max_column + 1):
        v = str(ws_q.cell(1, c).value or "").strip().upper()
        if v in target_headers:
          cols_map[v] = c

    datos_quin = {}
    for r in range(idx_enc + 1, ws_q.max_row + 1):
      t_val = ws_q.cell(r, cols_map["TARJETA"]).value
      k = clave(t_val)
      if not k:
        continue
      
      info = {}
      for h in target_headers:
        col_idx = cols_map.get(h)
        info[h] = ws_q.cell(r, col_idx).value if col_idx else None
      datos_quin[k] = info
      
    return datos_quin

  # Cargar datos dinámicos de las quincenas
  q1_dict = _cargar_datos_quincena(f_q1)
  q2_dict = _cargar_datos_quincena(f_q2) if f_q2 is not None else {}

  if hasattr(f_final, "seek"):
    f_final.seek(0)
  wb_orig = load_workbook(f_final, data_only=True)
  ws_orig = wb_orig.active

  if ws_orig.max_row < 2:
    raise ValueError("El archivo final no tiene filas de datos.")

  wb_new = Workbook()
  ws_new = wb_new.active
  ws_new.title = "FILTRADO_QUINCENAS"

  max_col_orig = ws_orig.max_column
  # 1. Copiar encabezado original (fila 1) con estilos
  for col_c in range(1, max_col_orig + 1):
    src_cell = ws_orig.cell(row=1, column=col_c)
    dest_cell = ws_new.cell(row=1, column=col_c, value=src_cell.value)
    if src_cell.has_style:
      dest_cell.font = copy.copy(src_cell.font)
      dest_cell.fill = copy.copy(src_cell.fill)
      dest_cell.border = copy.copy(src_cell.border)
      dest_cell.alignment = copy.copy(src_cell.alignment)
      dest_cell.number_format = src_cell.number_format

  # Escribir encabezados para las columnas adicionales a partir de la Z (columna 26)
  target_headers = ["PUESTO", "DESCR", "TIPOTRAB", "RAMA", "REGION"]
  for i, h in enumerate(target_headers):
    col_idx = 26 + i
    dest_cell = ws_new.cell(row=1, column=col_idx, value=h)
    dest_cell.font = Font(bold=True)
    dest_cell.alignment = Alignment(horizontal="center", vertical="center")

  filas_copiadas = 0
  total_filas = 0
  new_row_idx = 2

  # 2. Filtrar y copiar filas manteniendo formatos originales y agregando los datos dinámicos
  for r in range(2, ws_orig.max_row + 1):
    if all(es_vacio(ws_orig.cell(r, c).value) for c in range(1, max_col_orig + 1)):
      continue
    total_filas += 1

    # Columna B del archivo final (índice 2)
    val_b = ws_orig.cell(r, 2).value
    k_emp = clave(val_b)

    info_match = None
    if k_emp and k_emp in q1_dict:
      info_match = q1_dict[k_emp]
    elif k_emp and k_emp in q2_dict:
      info_match = q2_dict[k_emp]

    if info_match is not None:
      filas_copiadas += 1
      # Copiar las columnas originales del archivo final
      for col_c in range(1, max_col_orig + 1):
        src_cell = ws_orig.cell(row=r, column=col_c)
        dest_cell = ws_new.cell(row=new_row_idx, column=col_c, value=src_cell.value)
        if src_cell.has_style:
          dest_cell.font = copy.copy(src_cell.font)
          dest_cell.fill = copy.copy(src_cell.fill)
          dest_cell.border = copy.copy(src_cell.border)
          dest_cell.alignment = copy.copy(src_cell.alignment)
          dest_cell.number_format = copy.copy(src_cell.number_format)

      # Pegar PUESTO, DESCR, TIPOTRAB, RAMA y REGION a partir de la celda Z (columna 26)
      for i, h in enumerate(target_headers):
        ws_new.cell(row=new_row_idx, column=26 + i, value=info_match.get(h))

      new_row_idx += 1

  for col_letter, dim in ws_orig.column_dimensions.items():
    ws_new.column_dimensions[col_letter].width = dim.width

  ws_new.freeze_panes = "A2"
  buf = io.BytesIO()
  wb_new.save(buf)
  return buf.getvalue(), filas_copiadas, total_filas

# --------------------------------------------------------------------------
# Archivo de la quincena (lo usa la utilidad «Obtener personal EN FIRMA»)
# --------------------------------------------------------------------------
def leer_quincena(archivo):
  """Lee la quincena (Hoja1). Detecta el renglón de encabezados por la columna «tarjeta»."""
  filas = leer_hoja1(archivo)
  idx = None
  for i, r in enumerate(filas[:15]):
    if "tarjeta" in [str(c).strip().lower() for c in r if not es_vacio(c)]:
      idx = i
      break
  if idx is None:
    raise ValueError("En el archivo de la quincena no encontré la columna «tarjeta» "
            "(debe estar en los primeros renglones de la Hoja1).")
  enc = ["" if es_vacio(c) else str(c).strip().lower() for c in filas[idx]]
  pos = {n: enc.index(a) for n, alias, _ in COL_QUINCENA for a in alias if a in enc}
  faltan = [n for n, _, opc in COL_QUINCENA if not opc and n not in pos]
  if faltan:
    raise ValueError("A la quincena le faltan columnas: " + ", ".join(faltan))
  i_tar = enc.index("tarjeta")
  i_nom = enc.index("nombre") if "nombre" in enc else None

  registros = []
  for r in filas[idx + 1:]:
    t = clave(celda(r, i_tar))
    if not t:
      continue
    reg = {"TARJETA": t, "NOMBRE": valor(celda(r, i_nom)) if i_nom is not None else None}
    for n, i in pos.items():
      reg[n] = valor(celda(r, i))
    registros.append(reg)
  columnas = [n for n, _, _ in COL_QUINCENA if n in pos]
  return registros, columnas


def _norm_nombre(s) -> str:
  s = unicodedata.normalize("NFKD", str(s or "").upper())
  s = "".join(ch for ch in s if ch.isalpha() or ch == " ")
  return " ".join(sorted(s.split()))      # sin importar el orden nombre/apellidos


def elegir_fila(candidatas, nombre):
  """Si una tarjeta está repetida en la quincena, elige la fila con nombre más parecido."""
  if len(candidatas) == 1:
    return candidatas[0]
  objetivo = _norm_nombre(nombre)
  return max(candidatas,
       key=lambda c: SequenceMatcher(None, objetivo, _norm_nombre(c.get("NOMBRE"))).ratio())


def util_personal_en_firma(f_todos, f_quincena):
  """Macro 2: de todo el personal inactivo del sistema (export de MySQL, Hoja1, tarjeta en la

  columna A) obtiene a quienes aparecen en la quincena. Esa lista es la que hay
  que agregar a mano
  en el reporte de Inspector (con «mostrar empleados inactivos»).
  """
  filas = leer_tabla(f_todos)
  quin, cols_q = leer_quincena(f_quincena)
  indice = {}
  for q in quin:
    indice.setdefault(q["TARJETA"], []).append(q)
  cols = [c for c in COLS_EN_FIRMA if c in ("TARJETA", "NOMBRE") or c in cols_q]
  salida, vistos = [], set()
  for r in filas:
    t = clave(celda(r, 0))
    if not t or t in vistos or t not in indice:
      continue
    vistos.add(t)
    nombre = " ".join(
        str(x) for x in (celda(r, 2), celda(r, 3)) if not es_vacio(x)
    )
    q = elegir_fila(indice[t], nombre)
    salida.append([t if c == "TARJETA" else q.get(c) for c in cols])
  wb = Workbook()
  ws = wb.active
  ws.title = "EN FIRMA"
  ws.append(cols)
  for c in range(1, len(cols) + 1):
    x = ws.cell(row=1, column=c)
    x.font = Font(bold=True)
    x.alignment = Alignment(horizontal="center", vertical="center")
  for f in salida:
    ws.append(f)
  _autoajustar(ws)
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue(), pd.DataFrame(salida, columns=cols, dtype=object)


def util_buscar_gua(archivo, texto="gua", columna=5):
  """Macro 8: copia las columnas A:D de las filas cuya columna E contiene `texto`."""
  filas = leer_hoja1(archivo)
  wb = Workbook()
  ws = wb.active
  ws.title = "RESULTADO"
  if filas:
    ws.append([celda(filas[0], j) for j in range(4)])  # encabezados A:D
  encontradas = []
  for r in filas[1:]:
    v = celda(r, columna - 1)
    if not es_vacio(v) and texto.lower() in str(v).lower():
      fila = [celda(r, j) for j in range(4)]
      encontradas.append(fila)
      ws.append(fila)
  _autoajustar(ws)
  buf = io.BytesIO()
  wb.save(buf)
  cab = (
      _unicos([
          (("" if es_vacio(c) else str(c)) or f"col {i + 1}")
          for i, c in enumerate([celda(filas[0], j) for j in range(4)])
      ])
      if filas
      else None
  )
  return buf.getvalue(), pd.DataFrame(encontradas, columns=cab)


def util_formato_quincena(archivo):
  """Macro 9: reacomoda columnas. Origen A..U -> destino E..X:

  E=A, F=B, G=C, H=D, I:K=E:G, L:X=I:U (la columna H del origen se descarta).
  """
  filas = leer_hoja1(archivo)
  ultima = max(
      (i for i, r in enumerate(filas) if not es_vacio(celda(r, 0))), default=-1
  )
  wb = Workbook()
  ws = wb.active
  ws.title = "QUINCENA"
  for r in filas[: ultima + 1]:
    dest = [None] * 24
    dest[4], dest[5], dest[6], dest[7] = (
        celda(r, 0),
        celda(r, 1),
        celda(r, 2),
        celda(r, 3),
    )
    dest[8:11] = [celda(r, 4), celda(r, 5), celda(r, 6)]
    dest[11:24] = [celda(r, j) for j in range(8, 21)]
    ws.append(dest)
  buf = io.BytesIO()
  wb.save(buf)
  return buf.getvalue(), ultima + 1

def util_analizar_reporte(archivo):
  """Analiza el reporte para extraer el Turno (Columna T / índice 19),
  Descr (Columna AA / índice 26) y Tipo de Trabajador (Columna AB / índice 27).
  """
  if hasattr(archivo, "seek"):
    archivo.seek(0)
  df = pd.read_excel(archivo, sheet_name=0, dtype=object)
  if df.empty:
    raise ValueError("El archivo está vacío.")
  
  cols = list(df.columns)
  
  # Búsqueda por nombre de encabezado o por posición absoluta de columna Excel (T=19, AA=26, AB=27)
  c_turno = next((c for c in cols if str(c).strip().upper() == "TURNO"), cols[19] if len(cols) > 19 else None)
  c_descr = next((c for c in cols if str(c).strip().upper() in ("DESCR", "DESCRIPCION", "DESCRIPCIÓN")), cols[26] if len(cols) > 26 else None)
  c_tipotrab = next((c for c in cols if str(c).strip().upper() in ("TIPOTRAB", "TIPO DE TRABAJADOR", "TIPO_TRABAJADOR")), cols[27] if len(cols) > 27 else None)
  
  if not c_turno or not c_descr or not c_tipotrab:
    raise ValueError("No se pudieron localizar las columnas requeridas (Turno en T, Descr en AA, Tipo de Trabajador en AB).")
    
  return df, c_turno, c_descr, c_tipotrab


def exportar_por_turno(df, col_turno, turno_seleccionado):
  """Filtra el DataFrame por un turno específico (Columna T) y genera un Excel descargable."""
  df_filtrado = df[df[col_turno].astype(str).str.strip().str.upper() == str(turno_seleccionado).strip().upper()].copy()
  buf = io.BytesIO()
  with pd.ExcelWriter(buf, engine="openpyxl") as writer:
    df_filtrado.to_excel(writer, index=False, sheet_name="REPORTE_TURNO")
  return buf.getvalue(), len(df_filtrado)


def util_filtrar_por_tarjetas(f_final, texto_tarjetas):
  """Filtra el archivo final conservando formatos originales según una lista de
  tarjetas ingresadas como texto separadas por comas (buscando en la Columna A).
  """
  # 1. Limpiar y extraer la lista de tarjetas requeridas
  tarjetas_raw = [t.strip() for t in texto_tarjetas.split(",") if t.strip()]
  tarjetas_set = {clave(t) for t in tarjetas_raw if clave(t)}

  if not tarjetas_set:
    raise ValueError("No se ingresaron tarjetas válidas para filtrar.")

  if hasattr(f_final, "seek"):
    f_final.seek(0)
  wb_orig = load_workbook(f_final, data_only=True)
  ws_orig = wb_orig.active

  if ws_orig.max_row < 2:
    raise ValueError("El archivo final no tiene filas de datos.")

  wb_new = Workbook()
  ws_new = wb_new.active
  ws_new.title = "FILTRADO_TARJETAS"

  max_col_orig = ws_orig.max_column

  # 2. Copiar encabezado (fila 1) con sus estilos originales
  for col_c in range(1, max_col_orig + 1):
    src_cell = ws_orig.cell(row=1, column=col_c)
    dest_cell = ws_new.cell(row=1, column=col_c, value=src_cell.value)
    if src_cell.has_style:
      dest_cell.font = copy.copy(src_cell.font)
      dest_cell.fill = copy.copy(src_cell.fill)
      dest_cell.border = copy.copy(src_cell.border)
      dest_cell.alignment = copy.copy(src_cell.alignment)
      dest_cell.number_format = src_cell.number_format

  filas_encontradas = 0
  tarjetas_vistas = set()
  datos_vista = []
  
  # Usamos _unicos() para evitar el error de nombres de columnas duplicados en Streamlit/PyArrow
  raw_headers = [str(ws_orig.cell(1, c).value or f"Col {c}") for c in range(1, max_col_orig + 1)]
  headers = _unicos(raw_headers)
  
  new_row_idx = 2

  # 3. Filtrar las filas por coincidencia en la columna A (índice 1)
  for r in range(2, ws_orig.max_row + 1):
    if all(es_vacio(ws_orig.cell(r, c).value) for c in range(1, max_col_orig + 1)):
      continue

    # Leemos la celda de la Columna A (índice 1)
    val_a = ws_orig.cell(r, 1).value
    k_emp = clave(val_a)

    if k_emp and k_emp in tarjetas_set:
      filas_encontradas += 1
      tarjetas_vistas.add(k_emp)
      
      fila_valores = []
      for col_c in range(1, max_col_orig + 1):
        src_cell = ws_orig.cell(row=r, column=col_c)
        dest_cell = ws_new.cell(row=new_row_idx, column=col_c, value=src_cell.value)
        fila_valores.append(src_cell.value)
        
        if src_cell.has_style:
          dest_cell.font = copy.copy(src_cell.font)
          dest_cell.fill = copy.copy(src_cell.fill)
          dest_cell.border = copy.copy(src_cell.border)
          dest_cell.alignment = copy.copy(src_cell.alignment)
          dest_cell.number_format = src_cell.number_format
          
      datos_vista.append(fila_valores)
      new_row_idx += 1

  # Copiar dimensiones de columnas originales
  for col_letter, dim in ws_orig.column_dimensions.items():
    ws_new.column_dimensions[col_letter].width = dim.width

  ws_new.freeze_panes = "A2"
  buf = io.BytesIO()
  wb_new.save(buf)

  # Identificar tarjetas no localizadas en el archivo
  no_encontradas = [t for t in tarjetas_raw if clave(t) not in tarjetas_vistas]
  df_vista = pd.DataFrame(datos_vista, columns=headers, dtype=object) if datos_vista else pd.DataFrame()

  return buf.getvalue(), filas_encontradas, len(tarjetas_set), no_encontradas, df_vista
